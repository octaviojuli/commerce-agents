"""Serve isolated fictional ACME organizations for manual browser acceptance.

Uses only warehouse_test from the local private database configuration. No upstream
connector, real supplier identity, or source data is loaded. Model credentials are
loaded only when --live-model is explicitly selected; all business data stays fictional.
"""

import argparse
import asyncio
import json
import os
import sys
from contextlib import asynccontextmanager, suppress
from datetime import UTC, datetime, timedelta
from io import BytesIO
from pathlib import Path
from uuid import UUID, uuid4

import uvicorn
from openpyxl import Workbook
from sqlalchemy import text
from sqlalchemy.engine import make_url

from cloud_warehouse import auth, changes, documents, imports, inventory, pricing
from cloud_warehouse.admin import grant_auth, grant_runtime, migrate, onboard
from cloud_warehouse.api import create_app
from cloud_warehouse.assets import LocalObjectStore
from cloud_warehouse.integrations import SourceError
from cloud_warehouse.persistence import Principal, engine_for, transaction
from cloud_warehouse.pricing import CustomerIdentity

ROOT = Path(__file__).resolve().parents[2]
STATE = ROOT / ".warehouse"


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--seed-catalog", action="store_true")
    parser.add_argument("--seed-prices", action="store_true")
    parser.add_argument("--seed-display", action="store_true")
    parser.add_argument("--scripted-assistant", action="store_true")
    parser.add_argument(
        "--live-model",
        action="store_true",
        help="Use the configured model with fictional ACME data only",
    )
    parser.add_argument("--document-worker", action="store_true")
    parser.add_argument("--seed-attachments", action="store_true")
    parser.add_argument("--seed-sync", action="store_true")
    parser.add_argument("--seed-inventory", action="store_true")
    parser.add_argument("--seed-history", action="store_true")
    parser.add_argument("--seed-legacy-cases", action="store_true")
    args = parser.parse_args()
    if args.live_model and args.scripted_assistant:
        parser.error("Choose either the live model or scripted assistant")
    if args.live_model:
        from dotenv import dotenv_values

        model_settings = {**dotenv_values(ROOT / "examples/tour/.env"), **os.environ}
        for key in (
            "ANTHROPIC_API_KEY",
            "ANTHROPIC_AUTH_TOKEN",
            "ANTHROPIC_BASE_URL",
            "TOUR_MODEL",
        ):
            if model_settings.get(key):
                os.environ[key] = model_settings[key]
        if not (os.environ.get("ANTHROPIC_API_KEY") or os.environ.get("ANTHROPIC_AUTH_TOKEN")):
            parser.error("Live model credentials are not configured")
    os.umask(0o077)
    config = json.loads((STATE / "database.json").read_text())
    urls = {
        key: make_url(config[key])
        .set(database="warehouse_test")
        .render_as_string(hide_password=False)
        for key in ("admin_url", "runtime_url", "auth_url")
    }
    migrate(urls["admin_url"])
    admin = engine_for(urls["admin_url"])
    grant_runtime(admin, make_url(urls["runtime_url"]).username)
    grant_auth(admin, make_url(urls["auth_url"]).username)
    namespace = uuid4().hex[:10]
    first = onboard(
        admin, "ACME 云仓浏览器测试", "ACME 采购测试", f"web-sync-{namespace}@acme.example"
    )
    second = onboard(admin, "ACME 隔离组织", "ACME 隔离采购", f"web-other-{namespace}@acme.example")
    email, password = f"web-{namespace}@acme.example", "ACME-web-fixture-only-42"
    user = auth.create_user(
        admin,
        email,
        password,
        {
            UUID(first["supplier_id"]): ["supplier_admin"],
            UUID(second["supplier_id"]): ["auditor"],
            UUID(first["buyer_id"]): ["buyer_admin"],
            UUID(second["buyer_id"]): ["advisor"],
        },
    )
    source = uuid4()
    with admin.begin() as conn:
        conn.execute(
            text("""INSERT INTO supplier_connection(id,supplier_org_id,name,connector_type,credential_ref,capabilities)
          VALUES(:id,:org,'ACME Excel 交接','excel','none','{"inventory_owner":"warehouse"}'::jsonb)"""),
            {"id": source, "org": UUID(first["supplier_id"])},
        )
        conn.execute(
            text(
                "INSERT INTO distribution_grant(id,supplier_org_id,buyer_org_id,connection_id) VALUES(:id,:supplier,:buyer,:source)"
            ),
            {
                "id": uuid4(),
                "supplier": UUID(first["supplier_id"]),
                "buyer": UUID(first["buyer_id"]),
                "source": source,
            },
        )
    book = Workbook()
    sheet = book.active
    sheet.title = "团期导入"
    sheet.append(list(imports.HEADERS))
    for i in range(2):
        sheet.append(
            [
                f"ACME-R{i}",
                f"ACME 山海环线 {i + 1}",
                3,
                "ACME 城市",
                f"ACME-D{i}",
                "2026-10-01",
                "2026-10-03",
                20,
                2,
                1,
            ]
        )
    book.save(STATE / "浏览器测试团期.xlsx")
    sheet.cell(row=2, column=8, value="=1+1")
    book.save(STATE / "浏览器错误团期.xlsx")
    (STATE / "e2e-access.json").write_text(
        json.dumps(
            {
                "email": email,
                "password": password,
                "supplier_id": first["supplier_id"],
                "other_id": second["supplier_id"],
                "buyer_id": first["buyer_id"],
                "other_buyer_id": second["buyer_id"],
            },
            ensure_ascii=False,
        )
    )
    admin.dispose()
    runtime, authentication = engine_for(urls["runtime_url"]), engine_for(urls["auth_url"])
    try:
        if args.seed_display:
            from cloud_warehouse.catalog import synchronize
            from cloud_warehouse.integrations import CatalogBatch

            day = (datetime.now(UTC) + timedelta(days=30)).date()

            class DisplayCatalog:
                async def read_catalog(self, start=None, end=None):
                    return CatalogBatch(
                        [
                            {
                                "routeId": 80001,
                                "routeCode": "ACME-API",
                                "routeName": "ACME API 原始线路",
                                "days": 3,
                            }
                        ],
                        [
                            {
                                "periodId": 80002,
                                "routeId": 80001,
                                "periodCode": "ACME-API-D1",
                                "departDate": day.isoformat(),
                                "returnDate": (day + timedelta(days=2)).isoformat(),
                                "availableSeats": 21,
                            }
                        ],
                    )

            worker = Principal(UUID(first["worker_id"]), UUID(first["supplier_id"]))
            asyncio.run(
                synchronize(runtime, worker, UUID(first["connection_id"]), DisplayCatalog())
            )
            (STATE / "e2e-display.json").write_text(
                json.dumps(
                    {
                        "worker_id": first["worker_id"],
                        "connection_id": first["connection_id"],
                        "supplier_id": first["supplier_id"],
                        "day": day.isoformat(),
                    }
                )
            )
        if args.seed_attachments:
            from cloud_warehouse.catalog import synchronize
            from cloud_warehouse.integrations import CatalogBatch

            class FictionalCatalog:
                async def read_catalog(self, start=None, end=None):
                    return CatalogBatch(
                        [
                            {
                                "routeId": 70001,
                                "routeCode": "ACME-SOURCE",
                                "routeName": "ACME 来源线路",
                                "days": 3,
                                "routeAttachmentName": "ACME 来源行程.docx",
                                "routeAttachmentUrl": "https://files.acme.example/route.docx",
                            }
                        ],
                        [],
                    )

            worker = Principal(UUID(first["worker_id"]), UUID(first["supplier_id"]))
            asyncio.run(
                synchronize(runtime, worker, UUID(first["connection_id"]), FictionalCatalog())
            )
            # Deliberately start failed so the browser can exercise human retry. The
            # fixture later returns fictional bytes without contacting any file host.
            with transaction(runtime, worker) as conn:
                conn.execute(
                    text(
                        "UPDATE document_fetch SET status='failed',attempts=3,last_error='ATTACHMENT_NETWORK_FAILED' WHERE connection_id=:id"
                    ),
                    {"id": UUID(first["connection_id"])},
                )
        if args.seed_catalog or args.seed_prices or args.seed_inventory:
            actor = Principal(user, UUID(first["supplier_id"]))
            batch = imports.upload(
                runtime,
                actor,
                source,
                "ACME fixture.xlsx",
                (STATE / "浏览器测试团期.xlsx").read_bytes(),
            )
            proposal = imports.preview(runtime, actor, batch["id"])
            changes.approve(runtime, actor, UUID(proposal["id"]), proposal["payload_hash"])
            changes.apply(runtime, actor, UUID(proposal["id"]))
            if args.seed_inventory:
                with transaction(runtime, actor) as conn:
                    pool = conn.scalar(
                        text("""SELECT p.id FROM inventory_pool p
                        JOIN departure d ON d.id=p.departure_id
                        WHERE d.connection_id=:source AND d.code='ACME-D0'"""),
                        {"source": source},
                    )
                sale = None
                for index in range(506):
                    proposal = inventory.stage(
                        runtime,
                        actor,
                        inventory.InventoryCommand(
                            action="sale" if index == 0 else "adjust",
                            target_id=pool,
                            expected_version=index + 1,
                            quantity=3 if index == 0 else 1,
                            business_key="ACME-OLD-SALE"
                            if index == 0
                            else f"ACME-ADJUST-{index:04d}",
                            reason="ACME 虚构库存历史",
                        ),
                    )
                    changes.approve(runtime, actor, UUID(proposal["id"]), proposal["payload_hash"])
                    result = changes.apply(runtime, actor, UUID(proposal["id"]))
                    if index == 0:
                        sale = result["movement_id"]
                (STATE / "e2e-inventory.json").write_text(
                    json.dumps(
                        {
                            "pool_id": str(pool),
                            "sale_id": sale,
                            "rows": 507,
                            "available": 519,
                            "version": 507,
                        }
                    )
                )
            if args.seed_prices:
                with transaction(runtime, actor) as conn:
                    offers = (
                        conn.execute(
                            text("SELECT id FROM offer WHERE connection_id=:id"), {"id": source}
                        )
                        .scalars()
                        .all()
                    )
                for offer in offers:
                    proposal = pricing.propose(
                        runtime,
                        actor,
                        pricing.ContractProposal(
                            offer_id=offer,
                            buyer_org_id=UUID(first["buyer_id"]),
                            source_ref="ACME-TARIFF",
                            expected_version=0,
                            valid_from=datetime.now(UTC) - timedelta(days=1),
                            valid_until=datetime.now(UTC) + timedelta(days=30),
                            schedule=pricing.PriceSchedule(
                                currency="CNY",
                                market={"adult": "120.00", "child": "60.00"},
                                settlement={"adult": "100.00", "child": "50.00"},
                                fees_complete=True,
                                child_occupies_seat=True,
                                child_age_min=2,
                                child_age_max=12,
                                room_types=["双人标准间"],
                                included_items=["住宿", "交通"],
                            ),
                        ),
                    )
                    changes.approve(runtime, actor, UUID(proposal["id"]), proposal["payload_hash"])
                    changes.apply(runtime, actor, UUID(proposal["id"]))

        if args.seed_legacy_cases:
            from cloud_warehouse import legacy_cases
            from cloud_warehouse.catalog import synchronize
            from cloud_warehouse.integrations import CatalogBatch

            class LegacyCatalog:
                async def read_catalog(self, start=None, end=None):
                    return CatalogBatch(
                        [
                            {
                                "routeId": 36,
                                "routeCode": "ACME-LEGACY",
                                "routeName": "ACME 当前核验线路",
                                "days": 2,
                            }
                        ],
                        [],
                    )

            asyncio.run(
                synchronize(
                    runtime,
                    Principal(UUID(first["worker_id"]), UUID(first["supplier_id"])),
                    UUID(first["connection_id"]),
                    LegacyCatalog(),
                )
            )
            migration_admin = engine_for(urls["admin_url"])
            try:
                now = datetime.now(UTC)
                source_namespace = uuid4()
                identifiers = []
                for index in range(27):
                    versions = []
                    for version in range(1, 4):
                        versions.append(
                            {
                                "version": version,
                                "parent_version": None if version == 1 else 1,
                                "title": f"ACME 历史方案 v{version}",
                                "travel_dates": "历史出行日期",
                                "party": "2 人",
                                "days": [
                                    {
                                        "label": "第1天",
                                        "note": f"ACME 第 {version} 版行程安排",
                                        "request": version == 3,
                                    }
                                ],
                                "diff": []
                                if version == 1
                                else [
                                    {
                                        "index": 0,
                                        "kind": "changed",
                                        "fields": ["note"],
                                        "label": "第1天",
                                        "note": f"ACME 第 {version} 版行程安排",
                                    }
                                ],
                                "reference_price": {
                                    "tong_ye_adult": 100,
                                    "market_adult": 120,
                                    "party_total": None,
                                    "quote_source": "ACME 历史演示",
                                },
                                "share_token_hash": f"{index * 10 + version:064x}",
                                "created_at": now,
                            }
                        )
                    body = legacy_cases.Bundle(
                        source_namespace=source_namespace,
                        legacy_session_id=f"ACME-legacy-{index:02d}",
                        legacy_user_id="ACME 原顾问",
                        organization_id=UUID(first["buyer_id"]),
                        user_id=user,
                        connection_id=UUID(first["connection_id"]),
                        source_hash=f"{index:064x}",
                        created_at=now,
                        updated_at=now,
                        messages=[
                            {"role": "user", "text": f"ACME 历史案例 {index:02d}"},
                            {
                                "role": "assistant",
                                "text": "ACME 历史回答。<script>仅作为文字显示</script>",
                            },
                        ],
                        plans=[
                            {
                                "plan_id": f"PL-{index:08d}",
                                "route_id": 36,
                                "route_name": "ACME 历史线路",
                                "line_type": None,
                                "created_at": now,
                                "versions": versions,
                            }
                        ],
                    )
                    identifiers.append(
                        legacy_cases.import_case(
                            migration_admin,
                            body,
                            expected_database="warehouse_test",
                            note="ACME isolated browser fixture",
                            apply=True,
                        )["id"]
                    )
                (STATE / "e2e-legacy-cases.json").write_text(
                    json.dumps(
                        {
                            "ids": identifiers,
                            "grant_id": first["grant_id"],
                            "connection_id": first["connection_id"],
                            "user_id": str(user),
                            "organization_id": first["buyer_id"],
                        }
                    )
                )
            finally:
                migration_admin.dispose()

        class FictionalCustomerSource:
            async def resolve_customer(self, code):
                if code != "ACME-CODE":
                    raise SourceError("CUSTOMER_NOT_FOUND")
                return CustomerIdentity(
                    customer_id="ACME-UPSTREAM-001", code=code, name="ACME 采购测试"
                )

        if args.seed_history:
            actor = Principal(user, UUID(first["supplier_id"]))
            for index in range(131):
                history_book = Workbook()
                history_sheet = history_book.active
                history_sheet.title = "团期导入"
                history_sheet.append(list(imports.HEADERS))
                history_sheet.append(
                    [
                        f"ACME-H{index:03d}",
                        f"ACME 历史线路 {index:03d}",
                        1,
                        "ACME 城市",
                        f"ACME-HD{index:03d}",
                        "2026-10-01",
                        "2026-10-01",
                        20,
                        2,
                        1,
                    ]
                )
                body = BytesIO()
                history_book.save(body)
                batch = imports.upload(
                    runtime, actor, source, f"ACME 历史导入 {index:03d}.xlsx", body.getvalue()
                )
                proposal = imports.preview(runtime, actor, batch["id"])
                if index >= 31:
                    changes.discard(runtime, actor, UUID(proposal["id"]))
                if index == 0:
                    (STATE / "e2e-history.json").write_text(
                        json.dumps(
                            {
                                "oldest_batch": str(batch["id"]),
                                "oldest_change": proposal["id"],
                                "imports": 131,
                                "pending": 31,
                                "processed": 100,
                            }
                        )
                    )

        @asynccontextmanager
        async def fictional_source():
            yield FictionalCustomerSource()

        def scripted_agent(role, actor, buyer):
            # Fixed ACME responses exercise tools, not model comprehension.
            sys.path.insert(0, str(ROOT / "examples"))
            from cloud_warehouse import trip_brief
            from cloud_warehouse.advisor import WarehouseAdvisorBackend
            from commerce_common.testing import FakeClient, text_message, tool_calls_message
            from tour.api.warehouse_agent import build_agent

            if role != "advisor":
                raise ValueError("This fixture only scripts advisor conversations")
            backend = WarehouseAdvisorBackend(runtime, actor)

            class Scripted:
                agent = None

                async def stream_turn(self, messages, context, state):
                    saved = trip_brief.get(runtime, actor, UUID(context.session_id))

                    def call(name, arguments):
                        return tool_calls_message((name, arguments, str(uuid4())))

                    fields = {
                        k: {
                            "value": v,
                            "source": "inferred",
                            "hint": "ACME 脚本化测试值，待顾问确认",
                        }
                        for k, v in {
                            "destinations": ["ACME"],
                            "party_total": 2,
                            "adults": 2,
                            "children": 0,
                            "rooms": {"doubles": 1, "child_bed": None},
                            "window": {
                                "start": str(datetime.now(UTC).date()),
                                "end": str((datetime.now(UTC) + timedelta(days=365)).date()),
                            },
                        }.items()
                    }
                    replies = [
                        call(
                            "update_trip_brief",
                            {"expected_version": saved["version"], "fields": fields},
                        ),
                        call("search_routes", {}),
                        call(
                            "present_suggestions", {"suggestions": ["调整出行天数", "补充出发城市"]}
                        ),
                        text_message("ACME 浏览器验收：推断项需确认；仅查询，没有占位或下单。"),
                    ]
                    self.agent = build_agent(backend, client=FakeClient(replies))
                    async for item in self.agent.stream_turn(messages, context, state):
                        yield item

                async def aclose(self):
                    if self.agent and hasattr(self.agent, "aclose"):
                        await self.agent.aclose()

            return Scripted()

        print(
            "Fictional ACME browser fixture ready on 127.0.0.1:8006; access in private .warehouse/e2e-access.json"
        )
        object_store = LocalObjectStore(STATE / "e2e-objects")
        agent_factory = scripted_agent if args.scripted_assistant else None
        if args.live_model:
            sys.path.insert(0, str(ROOT / "examples"))
            from tour.api.warehouse_chat import factory

            agent_factory = factory(runtime, {})
        app = create_app(
            runtime,
            authentication,
            {UUID(first["connection_id"]): fictional_source},
            agent_factory=agent_factory,
            object_store=object_store,
        )
        if args.seed_sync:
            from cloud_warehouse import jobs
            from cloud_warehouse.integrations import CatalogBatch

            actor = Principal(UUID(first["worker_id"]), UUID(first["supplier_id"]))
            connection_id = UUID(first["connection_id"])

            class SyncCatalog:
                async def read_catalog(self, start=None, end=None):
                    await asyncio.sleep(1)
                    return CatalogBatch(
                        [
                            {
                                "routeId": 80001,
                                "routeCode": "ACME-SYNC",
                                "routeName": "ACME 同步线路",
                            }
                        ],
                        [
                            {
                                "periodId": 90000 + i,
                                "routeId": 0,
                                "periodCode": f"ACME-{i}",
                                "departDate": "2026-10-01",
                                "returnDate": "2026-10-03",
                            }
                            for i in range(511)
                        ],
                    )

            @asynccontextmanager
            async def sync_source():
                yield SyncCatalog()

            schedule = jobs.schedule(runtime, actor, connection_id)
            with transaction(runtime, actor) as conn:
                conn.execute(
                    text(
                        "UPDATE sync_job SET status='failed',attempts=5,last_error='SCAN_TIMEOUT' WHERE id=:id"
                    ),
                    {"id": schedule},
                )
                conn.execute(
                    text("""INSERT INTO sync_run(id,supplier_org_id,connection_id,status,actor_id)
                  VALUES(:id,:org,:source,'failed',:actor)"""),
                    [
                        {
                            "id": uuid4(),
                            "org": actor.organization_id,
                            "source": connection_id,
                            "actor": actor.user_id,
                        }
                        for _ in range(61)
                    ],
                )

            sync_base_lifespan = app.router.lifespan_context

            @asynccontextmanager
            async def sync_lifespan(application):
                async with sync_base_lifespan(application):

                    async def sync_work():
                        while True:
                            await jobs.run_once(runtime, actor, {connection_id: sync_source})
                            await asyncio.sleep(0.5)

                    task = asyncio.create_task(sync_work())
                    try:
                        yield
                    finally:
                        task.cancel()
                        with suppress(asyncio.CancelledError):
                            await task

            app.router.lifespan_context = sync_lifespan
        if args.document_worker or args.seed_attachments:
            sys.path.insert(0, str(ROOT / "examples"))
            from tour.api.warehouse_documents import parse

            base_lifespan = app.router.lifespan_context

            @asynccontextmanager
            async def documents_lifespan(application):
                async with base_lifespan(application):

                    async def work():
                        actor = Principal(UUID(first["worker_id"]), UUID(first["supplier_id"]))
                        while True:
                            if args.seed_attachments:
                                import io
                                import zipfile

                                from cloud_warehouse import document_fetches
                                from cloud_warehouse.remote_documents import Download

                                def download(*_):
                                    output = io.BytesIO()
                                    with zipfile.ZipFile(output, "w") as archive:
                                        archive.writestr(
                                            "[Content_Types].xml",
                                            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>',
                                        )
                                        lines = [
                                            line
                                            for day in range(1, 4)
                                            for line in (
                                                f"D{day} ACME 城市",
                                                "游览【ACME 景点】，含门票。",
                                            )
                                        ]
                                        archive.writestr(
                                            "word/document.xml",
                                            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>'
                                            + "".join(
                                                f"<w:p><w:r><w:t>{line}</w:t></w:r></w:p>"
                                                for line in lines
                                            )
                                            + "</w:body></w:document>",
                                        )
                                    return Download(output.getvalue(), "ACME-fixture")

                                await asyncio.to_thread(
                                    document_fetches.run_once,
                                    runtime,
                                    actor,
                                    UUID(first["connection_id"]),
                                    object_store,
                                    {"files.acme.example"},
                                    download,
                                )
                            await asyncio.to_thread(
                                documents.run_once, runtime, actor, object_store, parse
                            )
                            await asyncio.sleep(0.5)

                    task = asyncio.create_task(work())
                    try:
                        yield
                    finally:
                        task.cancel()
                        with suppress(asyncio.CancelledError):
                            await task

            app.router.lifespan_context = documents_lifespan
        uvicorn.run(app, host="127.0.0.1", port=8006, proxy_headers=False, access_log=False)

    finally:
        runtime.dispose()
        authentication.dispose()


if __name__ == "__main__":
    main()
