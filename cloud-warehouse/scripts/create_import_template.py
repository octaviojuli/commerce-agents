"""Generate the supplier's input-only workbook. All example values are fictional."""

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from cloud_warehouse.imports import HEADERS, parse_xlsx

TARGET = Path(__file__).resolve().parents[1] / "templates/供应商团期导入模板.xlsx"


def main():
    workbook = Workbook()
    instructions = workbook.active
    instructions.title = "填写说明"
    for row in [
        ["云仓供应商团期导入", "模板版本 2；继续接受旧版 10 列模板"],
        [
            "编辑位置",
            "团期导入页第 2 行起，蓝色字为可填写内容。示例均为虚构 ACME 数据，上传前请替换。",
        ],
        ["编号", "线路编号、团期编号为供应商固定业务编号；保留为文本，更新时不能重新编号。"],
        ["日期", "使用 YYYY-MM-DD；天数必须与含首尾的日期天数一致。"],
        ["库存口径", "总库存为容量，不是剩余库存；云仓可用量由总量减已售与停售得到。"],
        [
            "首次交接",
            "首次导入必须填写期初已售与期初停售，零也须填写。供应商须确认交接后所有外部销售及时登记云仓。",
        ],
        [
            "更新库存",
            "已有团期的期初已售、期初停售必须留空；更新文件只能调整总量，不能覆盖销售流水。",
        ],
        ["审批", "上传后先校验和预览，确认内容后再审批发布；预览后库存有变化须重新预览。"],
        ["云仓停售", "选填：停售 / 解除停售；留空保留现值，新团期默认未停售。停售不改动库存数字。"],
        [
            "云仓报名截止",
            "填写 YYYY-MM-DD HH:MM:SS（北京时间），也接受带时区的 ISO 时间；仅日期不接受。留空保留现值，填写“清空”明确移除。新团期留空表示未设置。",
        ],
        [
            "销售规则说明",
            "更改停售或截止时必须填写。提议人须有供应商管理员权限，或同时具有库存管理与产品编辑权限；须经管理员审批。",
        ],
        [
            "销售规则校验",
            "截止不得晚于团期业务时区的出发日结束；调整出发日期时也核对已有截止。解除本地限制不代表供应商已确认可报名。",
        ],
        ["输入限制", "仅支持本模板 XLSX；不接受公式、宏或外部链接。最多 5000 行、10 MB。"],
        [
            "重复与删除",
            "同文件重复上传返回原批次，已发布文件不能重复生效；缺失行不会自动删除在线团期。",
        ],
        ["价格", "此表只负责产品、团期与库存；报价需单独配置已授权的价格来源或协议价。"],
    ]:
        instructions.append(row)
    instructions.column_dimensions["A"].width = 18
    instructions.column_dimensions["B"].width = 105
    for row in instructions:
        for cell in row:
            cell.font = Font(name="Arial", size=11, color="263348")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        instructions.row_dimensions[row[0].row].height = 40
    sheet = workbook.create_sheet("团期导入")
    sheet.append(list(HEADERS))
    sheet.append(
        [
            "ACME-R001",
            "ACME 山水三日游",
            3,
            "示例出发城市",
            "ACME-D001",
            "2026-10-01",
            "2026-10-03",
            20,
            3,
            2,
            "停售",
            "2026-09-30 18:00:00",
            "ACME 示例，待供应商完成交接后再解除停售",
        ]
    )
    for cell in sheet[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="203F59")
    for cell in sheet[2]:
        cell.font = Font(name="Arial", color="0000FF")
        cell.fill = PatternFill("solid", fgColor="FFF6D6")
    for column in "ABCDEFGHIJKLM":
        sheet.column_dimensions[column].width = 22 if column != "B" else 36
    sheet.column_dimensions["L"].width = 34
    sheet.column_dimensions["M"].width = 55
    for column in ("A", "E", "F", "G", "L"):
        sheet[f"{column}2"].number_format = "@"
    counts = DataValidation(
        type="whole", operator="between", formula1="0", formula2="1000000", allow_blank=True
    )
    counts.error = "请填写非负整数"
    counts.showErrorMessage = True
    sheet.add_data_validation(counts)
    counts.add("H2:J5001")
    states = DataValidation(type="list", formula1='"停售,解除停售"', allow_blank=True)
    states.error = "请选择停售、解除停售，或留空保留现值"
    states.showErrorMessage = True
    states.errorStyle = "stop"
    sheet.add_data_validation(states)
    states.add("K2:K5001")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:M2"
    TARGET.parent.mkdir(exist_ok=True)
    workbook.save(TARGET)
    rows, errors = parse_xlsx(TARGET.read_bytes())
    assert len(rows) == 1 and not errors
    check = load_workbook(TARGET, data_only=False)
    assert not any(cell.data_type in {"f", "e"} for page in check for row in page for cell in row)
    check.close()
    print(TARGET)


if __name__ == "__main__":
    main()
