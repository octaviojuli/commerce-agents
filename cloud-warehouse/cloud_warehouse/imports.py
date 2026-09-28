"""Validate private XLSX files, preview versions, and atomically publish approved imports."""

import hashlib
from datetime import UTC, date, datetime
from io import BytesIO
from pathlib import PurePath
from uuid import UUID, uuid4, uuid5
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from defusedxml.common import DefusedXmlException
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import AwareDatetime, BaseModel, ConfigDict, Field, ValidationError, model_validator
from sqlalchemy import Connection, Engine, text

from . import changes, inventory, sales
from .changes import Conflict
from .integrations import canonical, fingerprint
from .persistence import Forbidden, Principal, require_role, transaction

LEGACY_HEADERS = {
    "线路编号": "product_code",
    "线路名称": "product_name",
    "天数": "days",
    "出发城市": "gateway",
    "团期编号": "departure_code",
    "出发日期": "depart_date",
    "返回日期": "return_date",
    "总库存": "total",
    "期初已售": "initial_sold",
    "期初停售": "initial_blocked",
}
HEADERS = {
    **LEGACY_HEADERS,
    "云仓停售": "sales_paused",
    "云仓报名截止（北京时间）": "booking_deadline",
    "销售规则说明": "sales_reason",
}
MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 5000


class ImportRow(BaseModel):
    model_config = ConfigDict(extra="forbid", str_strip_whitespace=True)
    product_code: str = Field(min_length=1, max_length=100)
    product_name: str = Field(min_length=1, max_length=300)
    days: int = Field(ge=1, le=365, strict=True)
    gateway: str = Field(min_length=1, max_length=100)
    departure_code: str = Field(min_length=1, max_length=100)
    depart_date: date
    return_date: date
    total: int = Field(ge=0, le=1_000_000, strict=True)
    initial_sold: int | None = Field(default=None, ge=0, le=1_000_000, strict=True)
    initial_blocked: int | None = Field(default=None, ge=0, le=1_000_000, strict=True)
    sales_paused: bool | None = Field(default=None, strict=True)
    booking_deadline: AwareDatetime | None = None
    clear_booking_deadline: bool = Field(default=False, strict=True)
    sales_reason: str | None = Field(default=None, min_length=1, max_length=500)

    @model_validator(mode="after")
    def chronology(self):
        if (
            self.return_date < self.depart_date
            or (self.return_date - self.depart_date).days + 1 != self.days
        ):
            raise ValueError("出发、返回日期与天数不一致")
        if (self.initial_sold or 0) + (self.initial_blocked or 0) > self.total:
            raise ValueError("总库存小于期初已售与停售之和")
        if self.clear_booking_deadline and self.booking_deadline is not None:
            raise ValueError("截止时间与清空标记冲突")
        if (
            self.sales_paused is not None
            or self.booking_deadline is not None
            or self.clear_booking_deadline
        ) and not self.sales_reason:
            raise ValueError("修改销售规则须填写说明")
        return self


def parse_xlsx(body: bytes) -> tuple[list[dict], list[dict]]:
    if len(body) > MAX_BYTES:
        return [], [{"row": 0, "code": "FILE_TOO_LARGE"}]
    try:
        with ZipFile(BytesIO(body)) as archive:
            entries = archive.infolist()
            if len(entries) > 1000 or sum(item.file_size for item in entries) > 50 * 1024 * 1024:
                return [], [{"row": 0, "code": "EXPANDED_FILE_TOO_LARGE"}]
            if any(
                "vbaproject" in item.filename.lower() or "externallinks/" in item.filename.lower()
                for item in entries
            ):
                return [], [{"row": 0, "code": "MACROS_OR_EXTERNAL_LINKS_FORBIDDEN"}]
        workbook = load_workbook(BytesIO(body), read_only=True, data_only=False, keep_links=False)
    except (
        BadZipFile,
        ValueError,
        KeyError,
        OSError,
        RuntimeError,
        ParseError,
        DefusedXmlException,
        InvalidFileException,
    ):
        return [], [{"row": 0, "code": "INVALID_XLSX"}]
    rows, errors = [], []
    try:
        if "团期导入" not in workbook.sheetnames:
            return [], [{"row": 0, "code": "MISSING_IMPORT_SHEET"}]
        sheet = workbook["团期导入"]
        if sheet.max_row > MAX_ROWS + 1 or sheet.max_column > len(HEADERS):
            return [], [{"row": 0, "code": "SHEET_LIMIT_EXCEEDED"}]
        iterator = sheet.iter_rows()
        first = next(iterator, ())
        names = [cell.value for cell in first]
        headers = HEADERS if names == list(HEADERS) else LEGACY_HEADERS
        if names != list(headers):
            return [], [{"row": 1, "code": "HEADERS_MUST_MATCH_TEMPLATE"}]
        seen, products = set(), {}
        for number, cells in enumerate(iterator, 2):
            if all(cell.value is None for cell in cells):
                continue
            if any(cell.data_type in {"f", "e"} for cell in cells):
                errors.append({"row": number, "code": "FORMULA_OR_ERROR_CELL"})
                continue
            raw = {field: cell.value for field, cell in zip(headers.values(), cells, strict=True)}
            if headers is HEADERS:
                state = raw["sales_paused"]
                if state not in (None, "", "停售", "解除停售"):
                    errors.append(
                        {"row": number, "field": "sales_paused", "code": "INVALID_SALES_STATE"}
                    )
                    continue
                raw["sales_paused"] = {"停售": True, "解除停售": False}.get(state)
                value = raw["booking_deadline"]
                if value == "清空":
                    raw["clear_booking_deadline"], raw["booking_deadline"] = True, None
                elif value in (None, ""):
                    raw["booking_deadline"] = None
                else:
                    try:
                        if isinstance(value, str) and ("T" in value or " " in value):
                            value = datetime.fromisoformat(value)
                        if not isinstance(value, datetime):
                            raise ValueError("timestamp required")
                        if value.tzinfo is None:
                            value = value.replace(tzinfo=sales.business_zone("Asia/Shanghai"))
                        raw["booking_deadline"] = value.astimezone(UTC)
                    except ValueError:
                        errors.append(
                            {
                                "row": number,
                                "field": "booking_deadline",
                                "code": "INVALID_BOOKING_DEADLINE",
                            }
                        )
                        continue
                if raw["sales_reason"] == "":
                    raw["sales_reason"] = None
            try:
                model = ImportRow.model_validate(raw)
            except ValidationError as error:
                errors.extend(
                    {"row": number, "field": ".".join(map(str, item["loc"])), "code": item["type"]}
                    for item in error.errors(include_input=False)
                )
                continue
            row = model.model_dump(mode="json")
            if row["departure_code"] in seen:
                errors.append({"row": number, "code": "DUPLICATE_DEPARTURE_CODE"})
            seen.add(row["departure_code"])
            product = _product(row)
            previous = products.setdefault(row["product_code"], product)
            if product != previous:
                errors.append({"row": number, "code": "CONFLICTING_PRODUCT_FIELDS"})
            rows.append(row)
        if not rows and not errors:
            errors.append({"row": 0, "code": "EMPTY_FILE"})
        return rows, errors
    except (ParseError, DefusedXmlException, ValueError, KeyError, OSError):
        return [], [{"row": 0, "code": "INVALID_SHEET_XML"}]
    finally:
        workbook.close()


def _product(row: dict) -> dict:
    return {key: row[key] for key in ("product_code", "product_name", "days", "gateway")}


def _connection(conn: Connection, connection_id: UUID, *, lock=False):
    row = (
        conn.execute(
            text(
                "SELECT id,version,connector_type,capabilities FROM supplier_connection WHERE id=:id AND supplier_org_id=warehouse_org_id() AND active"
                + (" FOR UPDATE" if lock else "")
            ),
            {"id": connection_id},
        )
        .mappings()
        .one_or_none()
    )
    if not row:
        raise Forbidden("供应商连接不存在或无权限")
    if (
        row["connector_type"] != "excel"
        or row["capabilities"].get("inventory_owner") != "warehouse"
    ):
        raise Conflict("Excel 只允许导入约定由云仓管理库存的供应商连接")
    return row


def _sales_rules(conn, batch):
    source = _connection(conn, batch["connection_id"])
    zone_name = source["capabilities"].get("business_timezone", sales.DEFAULT_BUSINESS_TIMEZONE)
    zone = sales.business_zone(zone_name)
    ids = [
        uuid5(batch["connection_id"], "departure:" + row["departure_code"]) for row in batch["rows"]
    ]
    current = {
        str(row["id"]): row
        for row in conn.execute(
            text("SELECT id,sales_paused,local_booking_deadline FROM departure WHERE id=ANY(:ids)"),
            {"ids": ids},
        ).mappings()
    }
    rules = []
    for raw, identity in zip(batch["rows"], ids, strict=True):
        row = ImportRow.model_validate(raw)
        before = current.get(str(identity), {})
        old_paused, old_deadline = (
            before.get("sales_paused", False),
            before.get("local_booking_deadline"),
        )
        paused = old_paused if row.sales_paused is None else row.sales_paused
        deadline = None if row.clear_booking_deadline else row.booking_deadline or old_deadline
        if deadline and deadline > sales.departure_day_end(row.depart_date, zone):
            raise Conflict(
                f"团期 {row.departure_code}：报名截止晚于出发业务日结束，请明确修改或清空"
            )
        if (paused, deadline) == (old_paused, old_deadline):
            continue
        rules.append(
            {
                "target_id": str(identity),
                "departure_code": row.departure_code,
                "before_paused": old_paused,
                "sales_paused": paused,
                "before_deadline": old_deadline.astimezone(UTC).isoformat()
                if old_deadline
                else None,
                "local_booking_deadline": deadline.astimezone(UTC).isoformat()
                if deadline
                else None,
                "reason": row.sales_reason,
                "business_timezone": zone_name,
            }
        )
    return rules, source["version"]


def upload(
    engine: Engine, actor: Principal, connection_id: UUID, file_name: str, body: bytes
) -> dict:
    if len(body) > MAX_BYTES:
        raise Conflict("文件超过 10 MB 限制")
    with transaction(engine, actor) as conn:
        require_role(conn, "supplier_admin", "inventory_manager")
        _connection(conn, connection_id)
    rows, errors = parse_xlsx(body)
    digest, batch_id = hashlib.sha256(body).hexdigest(), uuid4()
    with transaction(engine, actor) as conn:
        require_role(conn, "supplier_admin", "inventory_manager")
        _connection(conn, connection_id)
        conn.execute(
            text(
                "INSERT INTO import_batch(id,supplier_org_id,connection_id,file_hash,file_name,file_body,rows,errors,status,created_by) VALUES(:id,:org,:connection,:hash,:name,:body,CAST(:rows AS jsonb),CAST(:errors AS jsonb),:status,:actor) ON CONFLICT(connection_id,file_hash) DO NOTHING"
            ),
            {
                "id": batch_id,
                "org": actor.organization_id,
                "connection": connection_id,
                "hash": digest,
                "name": PurePath(file_name).name[:200],
                "body": body,
                "rows": canonical(rows),
                "errors": canonical(errors),
                "status": "invalid" if errors else "validated",
                "actor": actor.user_id,
            },
        )
        batch = (
            conn.execute(
                text(
                    "SELECT id,status,errors,version,published_change_id FROM import_batch WHERE connection_id=:connection AND file_hash=:hash"
                ),
                {"connection": connection_id, "hash": digest},
            )
            .mappings()
            .one()
        )
        return {**dict(batch), "row_count": len(rows), "duplicate": batch["id"] != batch_id}


def _batch(conn, batch_id: UUID, *, lock=False):
    row = (
        conn.execute(
            text(
                "SELECT id,connection_id,file_hash,rows,errors,status,version FROM import_batch WHERE id=:id"
                + (" FOR UPDATE" if lock else "")
            ),
            {"id": batch_id},
        )
        .mappings()
        .one_or_none()
    )
    if not row:
        raise Forbidden("导入批次不存在或无权限")
    if row["status"] != "validated" or row["errors"]:
        raise Conflict("只允许发布校验通过且尚未发布的文件")
    _connection(conn, row["connection_id"], lock=lock)
    return row


def _versions(conn: Connection, batch, *, lock=False) -> dict:
    connection = batch["connection_id"]
    products = {uuid5(connection, "route:" + row["product_code"]) for row in batch["rows"]}
    departures = {uuid5(connection, "departure:" + row["departure_code"]) for row in batch["rows"]}
    result = {}
    for table, ids, key in (
        ("supplier_product", products, "id"),
        ("departure", departures, "id"),
        ("inventory_pool", departures, "departure_id"),
    ):
        values = {str(identity): 0 for identity in ids}
        # Table and column names come from the fixed list above, never request strings.
        current = conn.execute(
            text(
                f"SELECT {key},version FROM {table} WHERE {key}=ANY(:ids) ORDER BY {key}"
                + (" FOR UPDATE" if lock else "")
            ),
            {"ids": sorted(ids)},
        ).all()
        values.update({str(identity): version for identity, version in current})
        result[table] = values
    return result


def preview(engine: Engine, actor: Principal, batch_id: UUID) -> dict:
    with transaction(engine, actor) as conn:
        require_role(conn, "supplier_admin", "inventory_manager")
        batch = _batch(conn, batch_id)
        rules, source_version = _sales_rules(conn, batch)
        payload = {
            "batch_id": str(batch_id),
            "file_hash": batch["file_hash"],
            "version": batch["version"],
            "expected": _versions(conn, batch),
            "rows": batch["rows"],
            "sales_rules": rules,
            "source_version": source_version,
        }
    return changes.stage(engine, actor, "excel_import", payload)


def validate_change(
    conn: Connection, actor: Principal, payload: dict, *, lock=False
) -> tuple[UUID, int]:
    batch_id = UUID(payload["batch_id"])
    batch = _batch(conn, batch_id, lock=lock)
    if (
        batch["file_hash"] != payload["file_hash"]
        or batch["version"] != payload["version"]
        or batch["rows"] != payload["rows"]
    ):
        raise Conflict("导入内容与已校验文件不一致")
    if payload["expected"] != _versions(conn, batch, lock=lock):
        raise Conflict("预览后产品、团期或库存已变更，请重新生成预览")
    rules, source_version = _sales_rules(conn, batch)
    if (
        rules != payload.get("sales_rules", [])
        or payload.get("source_version", source_version) != source_version
    ):
        raise Conflict("预览后销售规则或供应源已变更，请重新生成预览")
    if rules:
        require_role(conn, "supplier_admin", "product_editor")
    for row in batch["rows"]:
        departure_id = uuid5(batch["connection_id"], "departure:" + row["departure_code"])
        if payload["expected"]["inventory_pool"][str(departure_id)] == 0:
            if row["initial_sold"] is None or row["initial_blocked"] is None:
                raise Conflict("首次交接必须明确填写期初已售和停售，零也必须填写")
        elif row["initial_sold"] is not None or row["initial_blocked"] is not None:
            raise Conflict("已有库存的期初列必须留空；销售与停售须通过登记、冲正或调整处理")
    return batch_id, batch["version"]


def apply_command(conn: Connection, actor: Principal, change_id: UUID, payload: dict) -> dict:
    batch = _batch(conn, UUID(payload["batch_id"]), lock=True)
    connection, now, run_id = batch["connection_id"], datetime.now(UTC), uuid4()
    products = {row["product_code"]: _product(row) for row in batch["rows"]}
    shared = {
        "org": actor.organization_id,
        "connection": connection,
        "now": now,
        "run": run_id,
        "actor": actor.user_id,
    }
    conn.execute(
        text(
            "INSERT INTO sync_run(id,supplier_org_id,connection_id,status,actor_id,product_count,departure_count,source_departure_count,source_hash,completed_at) VALUES(:run,:org,:connection,'published',:actor,:products,:departures,:departures,:hash,:now)"
        ),
        {
            **shared,
            "products": len(products),
            "departures": len(batch["rows"]),
            "hash": batch["file_hash"],
        },
    )
    for code, product in products.items():
        product_id = uuid5(connection, "route:" + code)
        conn.execute(
            text("""INSERT INTO supplier_product(id,supplier_org_id,connection_id,external_id,code,name,days,gateway,source,source_hash,status,observed_at,sync_run_id)
            VALUES(:id,:org,:connection,:code,:code,:name,:days,:gateway,CAST(:source AS jsonb),:hash,'published',:now,:run)
            ON CONFLICT(connection_id,external_id) DO UPDATE SET name=EXCLUDED.name,days=EXCLUDED.days,gateway=EXCLUDED.gateway,
            source=EXCLUDED.source,source_hash=EXCLUDED.source_hash,observed_at=EXCLUDED.observed_at,sync_run_id=EXCLUDED.sync_run_id,
            version=supplier_product.version+CASE WHEN supplier_product.source_hash<>EXCLUDED.source_hash
              OR ROW(supplier_product.name,supplier_product.days,supplier_product.gateway)
                 IS DISTINCT FROM ROW(EXCLUDED.name,EXCLUDED.days,EXCLUDED.gateway) THEN 1 ELSE 0 END"""),
            {
                **shared,
                "id": product_id,
                "code": code,
                "name": product["product_name"],
                "days": product["days"],
                "gateway": product["gateway"],
                "source": canonical(product),
                "hash": fingerprint(product),
            },
        )
    inventory_results = []
    sales_rules = {row["target_id"]: row for row in payload.get("sales_rules", [])}
    for row in batch["rows"]:
        departure_id = uuid5(connection, "departure:" + row["departure_code"])
        source = {
            key: value
            for key, value in row.items()
            if key
            not in {
                "total",
                "initial_sold",
                "initial_blocked",
                "sales_paused",
                "booking_deadline",
                "clear_booking_deadline",
                "sales_reason",
            }
        }
        rule = sales_rules.get(str(departure_id))
        conn.execute(
            text("""INSERT INTO departure(id,product_id,supplier_org_id,connection_id,external_id,code,depart_date,return_date,source,source_hash,observed_at,availability_expires_at,sync_run_id,sales_paused,local_booking_deadline)
            VALUES(:id,:product,:org,:connection,:code,:code,:start,:end,CAST(:source AS jsonb),:hash,:now,:now,:run,:paused,:deadline)
            ON CONFLICT(connection_id,external_id) DO UPDATE SET product_id=EXCLUDED.product_id,depart_date=EXCLUDED.depart_date,return_date=EXCLUDED.return_date,
            source=EXCLUDED.source,source_hash=EXCLUDED.source_hash,observed_at=EXCLUDED.observed_at,sync_run_id=EXCLUDED.sync_run_id,
            sales_paused=CASE WHEN :sales_change THEN EXCLUDED.sales_paused ELSE departure.sales_paused END,
            local_booking_deadline=CASE WHEN :sales_change THEN EXCLUDED.local_booking_deadline ELSE departure.local_booking_deadline END,
            version=departure.version+CASE WHEN departure.source_hash<>EXCLUDED.source_hash OR :sales_change THEN 1 ELSE 0 END"""),
            {
                **shared,
                "id": departure_id,
                "product": uuid5(connection, "route:" + row["product_code"]),
                "code": row["departure_code"],
                "start": row["depart_date"],
                "end": row["return_date"],
                "source": canonical(source),
                "hash": fingerprint(source),
                "sales_change": rule is not None,
                "paused": rule["sales_paused"] if rule else False,
                "deadline": datetime.fromisoformat(rule["local_booking_deadline"])
                if rule and rule["local_booking_deadline"]
                else None,
            },
        )
        pool = (
            conn.execute(
                text(
                    "SELECT id,total,version FROM inventory_pool WHERE departure_id=:id FOR UPDATE"
                ),
                {"id": departure_id},
            )
            .mappings()
            .one_or_none()
        )
        if pool and pool["total"] == row["total"]:
            continue
        command = inventory.InventoryCommand(
            action="adjust" if pool else "open",
            target_id=pool["id"] if pool else departure_id,
            expected_version=pool["version"] if pool else 0,
            quantity=row["total"] - pool["total"] if pool else row["total"],
            business_key="import:" + str(batch["id"]),
            reason="经复核的 Excel 库存交接或总量调整",
            initial_sold=row["initial_sold"] or 0,
            initial_blocked=row["initial_blocked"] or 0,
        )
        inventory_results.append(
            inventory.apply_command(conn, actor, change_id, command.model_dump(mode="json"))
        )
    conn.execute(
        text(
            "UPDATE import_batch SET status='published',version=version+1,published_change_id=:change WHERE id=:id"
        ),
        {"change": change_id, "id": batch["id"]},
    )
    return {
        "batch_id": str(batch["id"]),
        "products": len(products),
        "departures": len(batch["rows"]),
        "inventory_movements": len(inventory_results),
    }
